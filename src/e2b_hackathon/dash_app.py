import os
import sys
import json
import base64
import traceback
import re
from datetime import datetime
from typing import List, Dict, Any, Optional
import io

import dash
from dash import html, dcc, callback, Input, Output, State
import dash_bootstrap_components as dbc  # type: ignore

# import plotly.graph_objects as go  # type: ignore
import plotly.express as px  # type: ignore
import structlog
import pandas as pd
import numpy as np

# For handling different file types
from openai import OpenAI

from e2b_hackathon.pickle_analyzer import analyze_pickle_files

# Import the execute_code_with_data function from our new module
from e2b_hackathon.execute_code_with_data import execute_code_with_data

# Initialize the Dash app
app = dash.Dash(
    __name__,
    external_stylesheets=[dbc.themes.BOOTSTRAP],
    suppress_callback_exceptions=True,
    title="Datalyzer",
)
log = structlog.get_logger()
SEED = 42

# Define the navbar with logo
navbar = dbc.Navbar(
    dbc.Container(
        [
            html.A(
                dbc.Row(
                    [
                        dbc.Col(html.Img(src="/assets/D2MA-logo.svg", height="40px")),
                    ],
                    align="center",
                ),
                href="/",
            ),
        ],
        fluid=True,
    ),
    color="rgb(235, 248, 235)",  # Light pastel green
    dark=False,
    className="mb-4",
    fixed="top",  # Makes navbar sticky at top
    style={"width": "100%"},  # Ensures full width
)

# Initialize OpenAI client (will use API key from environment)
try:
    openai_client: Optional[OpenAI] = OpenAI()
except Exception as e:
    log.error("Failed to initialize OpenAI client", error=str(e))
    openai_client = None  # type: ignore

# Define application layout
layout = html.Div(
    [
        navbar,  # Add the navbar at the top
        dbc.Container(
            [
                # Add margin-top to account for fixed navbar
                html.Div(style={"margin-top": "80px"}),  # Spacer for fixed navbar
                dbc.Row(
                    [
                        dbc.Col(
                            [
                                html.H1(
                                    "Enhanced Data Analysis Platform",
                                    className="text-center my-4",
                                ),
                                html.P(
                                    "Upload data files (Pickle, CSV, Parquet) and context files (Text, PDF, Word) for comprehensive analysis",
                                    className="text-center lead mb-4",
                                ),
                            ],
                            width=12,
                        )
                    ]
                ),
                dbc.Row(
                    [
                        dbc.Col(
                            [
                                dbc.Card(
                                    [
                                        dbc.CardBody(
                                            [
                                                html.H5(
                                                    "Unified File Upload",
                                                    className="card-title",
                                                ),
                                                html.P(
                                                    "Upload any files (.pkl, .pickle, .csv, .parquet, .txt, .pdf, .docx, .json, .geojson) for analysis. Files will be automatically categorized based on extension.",
                                                    className="card-text text-muted",
                                                ),
                                                dcc.Upload(
                                                    id="unified-upload",
                                                    children=html.Div(
                                                        [
                                                            html.I(
                                                                className="fas fa-upload me-2"
                                                            ),
                                                            "Drag and Drop or ",
                                                            html.A("Select Files"),
                                                        ]
                                                    ),
                                                    style={
                                                        "width": "100%",
                                                        "height": "60px",
                                                        "lineHeight": "60px",
                                                        "borderWidth": "1px",
                                                        "borderStyle": "dashed",
                                                        "borderRadius": "5px",
                                                        "textAlign": "center",
                                                        "margin": "10px 0",
                                                    },
                                                    multiple=True,
                                                ),
                                                html.P(
                                                    "Files will be automatically analyzed upon upload",
                                                    className="text-muted small mt-2",
                                                ),
                                            ]
                                        )
                                    ],
                                    className="mb-4",
                                ),
                            ],
                            width=12,
                        ),
                    ]
                ),
                # Loading spinner for when analysis is running
                dbc.Row(
                    [
                        dbc.Col(
                            [
                                dcc.Loading(
                                    id="loading-analysis",
                                    type="circle",
                                    children=html.Div(id="loading-output"),
                                ),
                            ],
                            width=12,
                        )
                    ]
                ),
                # Display files currently being analyzed
                dbc.Row([dbc.Col([html.Div(id="files-being-analyzed")], width=12)]),
                # Error messages
                dbc.Row([dbc.Col([html.Div(id="error-message")], width=12)]),
                # Results section
                dbc.Row(
                    [
                        dbc.Col(
                            [
                                html.Div(id="data-output", className="mt-4"),
                            ],
                            width=12,
                        )
                    ]
                ),
                # Context files section
                dbc.Row(
                    [
                        dbc.Col(
                            [
                                html.Div(id="context-output", className="mt-4"),
                            ],
                            width=12,
                        )
                    ]
                ),
                # Correlation analysis section
                dbc.Row(
                    [
                        dbc.Col(
                            [
                                html.H3("Correlation Analysis", className="mt-4 mb-3"),
                                html.Button(
                                    "Analyze Correlations",
                                    id="analyze-correlations-button",
                                    className="btn btn-primary mb-3",
                                ),
                                dcc.Loading(
                                    id="loading-correlation",
                                    type="circle",
                                    children=html.Div(id="correlation-loading-output"),
                                ),
                                html.Div(id="correlation-output"),
                            ],
                            width=12,
                        )
                    ]
                ),
                # AI Analysis Plan section
                dbc.Row(
                    [
                        dbc.Col(
                            [
                                html.H3("AI Analysis Plan", className="mt-4 mb-3"),
                                html.Button(
                                    "Generate Analysis Plan",
                                    id="generate-plan-button",
                                    className="btn btn-primary mb-3",
                                ),
                                dcc.Loading(
                                    id="loading-plan",
                                    type="circle",
                                    children=html.Div(id="plan-loading-output"),
                                ),
                                html.Div(id="ai-plan-output"),
                            ],
                            width=12,
                        )
                    ]
                ),
                # AI Analysis Results section
                dbc.Row(
                    [
                        dbc.Col(
                            [
                                html.H3("AI Analysis Results", className="mt-4 mb-3"),
                                html.Div(id="ai-results-output"),
                            ],
                            width=12,
                        )
                    ]
                ),
                # Store components for data
                dcc.Store(id="data-store", storage_type="memory"),
                dcc.Store(id="context-store", storage_type="memory"),
                dcc.Store(id="correlation-store", storage_type="memory"),
                dcc.Store(id="ai-plan-store", storage_type="memory"),
                # Footer
                dbc.Row(
                    [
                        dbc.Col(
                            [
                                html.Hr(),
                                html.P(
                                    "Powered by E2B Sandbox and OpenAI - A secure environment for comprehensive data analysis",
                                    className="text-center text-muted small",
                                ),
                            ],
                            width=12,
                        )
                    ],
                    className="mt-5",
                ),
            ],
            fluid=True,
        ),
    ],
    style={"padding": "20px"},
)

# Set the layout
app.layout = layout  # type: ignore


def save_uploaded_files(content, filename, file_type="data"):
    """
    Save uploaded files to temporary directory

    Args:
        content: File content
        filename: File name
        file_type: Type of file ('data' or 'context')

    Returns:
        Path to the saved file
    """
    valid_extensions = {
        "data": [
            ".pkl",
            ".pickle",
            ".csv",
            ".parquet",
            ".xls",
            ".xlsx",
            ".json",
            ".geojson",
        ],
        "context": [".txt", ".pdf", ".docx", ".doc", ".md"],
    }

    # Check extension
    _, ext = os.path.splitext(filename)
    ext = ext.lower()
    if ext not in valid_extensions[file_type]:
        raise ValueError(f"Invalid file extension for {file_type} file: {ext}")

    # Create temporary directory if it doesn't exist
    temp_dir = os.path.join(os.getcwd(), "temp_files")
    os.makedirs(temp_dir, exist_ok=True)

    # Decode and save file
    content_type, content_string = content.split(",")
    decoded = base64.b64decode(content_string)

    # Create a clean filename
    clean_filename = "".join(
        [c if c.isalnum() or c in "._- " else "_" for c in filename]
    )
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    file_path = os.path.join(temp_dir, f"{timestamp}_{clean_filename}")

    # Save the file
    with open(file_path, "wb") as f:
        f.write(decoded)

    return file_path


def parse_analysis_results(analysis_text: str) -> Dict[str, Dict[str, Any]]:
    """
    Parse the raw analysis text to extract structured information

    Args:
        analysis_text: The raw text output from pickle_analyzer

    Returns:
        Dictionary with parsed analysis results
    """
    # Debug log the analysis text
    log.debug(
        "Parsing analysis results",
        text_length=len(analysis_text),
        text_sample=analysis_text[:200] + ("..." if len(analysis_text) > 200 else ""),
    )

    # Split into sections by file
    sections = analysis_text.split("=" * 50)
    sections = [s.strip() for s in sections if s.strip()]

    log.debug(f"Found {len(sections)} sections in analysis text")

    results = {}

    for i, section in enumerate(sections):
        lines = section.split("\n")
        if len(lines) < 2:
            log.warning(f"Section {i} too short to parse", lines_count=len(lines))
            continue

        # Extract filename from the first line
        first_line = lines[0]
        log.debug(f"Processing section {i}, first line: {first_line}")

        # Check for the expected "Analysis of [filename]:" pattern
        if "Analysis of " in first_line and ":" in first_line:
            file_name = first_line.replace("Analysis of ", "").replace(":", "").strip()
            log.debug(f"Extracted file name: '{file_name}'")
        else:
            log.warning(
                f"Couldn't extract file name from section {i}", first_line=first_line
            )
            # Try a more flexible pattern
            file_name = f"unknown_file_{i}"
            for line in lines[:3]:  # Check first few lines
                if ".pkl" in line or ".pickle" in line:
                    potential_name = line.split("/")[-1].split()[0]
                    if potential_name.endswith(".pkl") or potential_name.endswith(
                        ".pickle"
                    ):
                        file_name = potential_name
                        log.debug(f"Found potential file name: {file_name}")
                        break

        # Extract other information
        file_info = {}
        current_key = None
        current_value = []

        for line in lines[1:]:
            line = line.strip()
            if not line:
                continue

            # If it's a new key-value pair
            if ": " in line and not line.startswith(" "):
                # Save the previous key-value pair if it exists
                if current_key:
                    file_info[current_key] = (
                        "\n".join(current_value)
                        if len(current_value) > 1
                        else current_value[0]
                    )
                    current_value = []

                # Start a new key-value pair
                current_key, value = line.split(": ", 1)
                current_value.append(value)
            else:
                # Continue with the current value
                if current_key:
                    current_value.append(line)

        # Save the last key-value pair
        if current_key:
            file_info[current_key] = (
                "\n".join(current_value) if len(current_value) > 1 else current_value[0]
            )

        results[file_name] = file_info
        log.debug(f"Added results for '{file_name}' with {len(file_info)} properties")

    log.info(
        f"Finished parsing results for {len(results)} files",
        file_names=list(results.keys()),
    )
    return results


def create_file_summary_card(file_name: str, file_info: Dict[str, Any]) -> dbc.Card:
    """
    Create a card with a summary of the file analysis

    Args:
        file_name: Name of the file
        file_info: Dictionary with file information

    Returns:
        A Dash Bootstrap Components Card
    """
    # Extract file type for additional component generation
    file_type = file_info.get("Type", "Unknown")

    # Create list group items dynamically from file_info
    list_group_items = []
    for key, value in file_info.items():
        # Skip complex nested data that would be better displayed in the additional info section
        if isinstance(value, (dict, list, tuple)) or key in [
            "sample_keys",
            "sample_elements",
            "sample_data",
            "columns",
            "dtypes",
        ]:
            continue

        # Format the value as string if it's not already
        if not isinstance(value, str):
            value = str(value)

        list_group_items.append(
            dbc.ListGroupItem([html.Strong(f"{key}: "), html.Span(value)])
        )

    # Create card content
    card_content = [
        dbc.CardHeader(html.H5(file_name, className="mb-0")),
        dbc.CardBody(
            [
                html.H6("Basic Information", className="card-subtitle mb-2 text-muted"),
                dbc.ListGroup(list_group_items, className="mb-3"),
                # Additional information based on file type
                get_additional_info_component(file_type, file_info),
            ]
        ),
    ]

    return dbc.Card(card_content, className="mb-4")


def get_additional_info_component(
    file_type: str, file_info: Dict[str, Any]
) -> html.Div:
    """
    Create appropriate components to display additional information based on file type

    Args:
        file_type: Type of the file (dict, list, DataFrame, etc.)
        file_info: Dictionary with file information

    Returns:
        Dash HTML component
    """
    components = []

    if "Pickle File" in file_type:
        # Information for pickle file
        analyzed_objects = file_info.get("Analyzed Objects", [])
        if analyzed_objects:
            # Display summary of analyzed objects
            object_count = len(analyzed_objects)
            object_types = file_info.get("Object Type Summary", {})

            components.extend(
                [
                    html.H6("Pickle Analysis Summary", className="mt-3"),
                    dbc.ListGroup(
                        [
                            dbc.ListGroupItem(
                                [
                                    html.Strong("Objects Found: "),
                                    html.Span(str(object_count)),
                                ]
                            ),
                            dbc.ListGroupItem(
                                [
                                    html.Strong("Object Types: "),
                                    html.Span(
                                        ", ".join(
                                            [
                                                f"{k} ({v})"
                                                for k, v in object_types.items()
                                            ]
                                        )
                                    ),
                                ]
                            ),
                        ],
                        className="mb-3",
                    ),
                ]
            )

            # Create cards for each object - directly visible without collapsible panels
            object_cards = []
            for i, obj in enumerate(analyzed_objects):
                obj_type = obj.get("Object Type", "Unknown")

                # Create list items for each property
                obj_details = []
                for key, value in obj.items():
                    if key not in ["Object Index"]:
                        obj_details.append(
                            dbc.ListGroupItem(
                                [html.Strong(f"{key}: "), html.Span(str(value))]
                            )
                        )

                # Create card for this object - directly visible
                object_cards.append(
                    dbc.Card(
                        [
                            dbc.CardHeader(
                                html.H6(
                                    f"Object #{i + 1}: {obj_type}", className="mb-0"
                                )
                            ),
                            dbc.CardBody([dbc.ListGroup(obj_details)]),
                        ],
                        className="mb-3",
                    )
                )

            components.append(
                html.Div(
                    [
                        html.H6("Analyzed Objects", className="mt-3"),
                        html.Div(object_cards),
                    ]
                )
            )
        else:
            # Show raw output if no objects were parsed
            raw_output = file_info.get("Raw Output", "No output available")
            components.append(
                html.Div(
                    [
                        html.H6("Raw Analysis Output", className="mt-3"),
                        dbc.Card(
                            dbc.CardBody(
                                html.Pre(
                                    raw_output,
                                    style={"maxHeight": "300px", "overflowY": "auto"},
                                )
                            ),
                            className="mb-3",
                        ),
                    ]
                )
            )
    elif file_type == "dict":
        # Information for dictionary
        keys_count = file_info.get("Number of keys", "Unknown")
        key_types = file_info.get("Key types", "Unknown")
        value_types = file_info.get("Value types", "Unknown")
        sample_keys = file_info.get("Sample keys", "Unknown")

        components.extend(
            [
                html.H6("Dictionary Details", className="mt-3"),
                dbc.ListGroup(
                    [
                        dbc.ListGroupItem(
                            [html.Strong("Number of keys: "), html.Span(keys_count)]
                        ),
                        dbc.ListGroupItem(
                            [html.Strong("Key types: "), html.Span(key_types)]
                        ),
                        dbc.ListGroupItem(
                            [html.Strong("Value types: "), html.Span(value_types)]
                        ),
                        dbc.ListGroupItem(
                            [html.Strong("Sample keys: "), html.Span(str(sample_keys))]
                        ),
                    ],
                    className="mb-3",
                ),
            ]
        )

        # Create a pie chart of value types if available
        if isinstance(value_types, str) and "," in value_types:
            try:
                value_type_list = [vt.strip() for vt in value_types.split(",")]
                value_counts = {vt: 1 for vt in value_type_list}

                fig = px.pie(
                    names=list(value_counts.keys()),
                    values=list(value_counts.values()),
                    title="Value Types Distribution",
                )

                components.append(dcc.Graph(figure=fig, className="mt-3"))
            except Exception:
                pass

    elif file_type == "list" or file_type == "tuple":
        # Information for list or tuple
        length = file_info.get("Length", "Unknown")
        element_types = file_info.get("Element types", "Unknown")
        sample_elements = file_info.get("Sample elements", "Unknown")

        components.extend(
            [
                html.H6(f"{file_type.capitalize()} Details", className="mt-3"),
                dbc.ListGroup(
                    [
                        dbc.ListGroupItem([html.Strong("Length: "), html.Span(length)]),
                        dbc.ListGroupItem(
                            [html.Strong("Element types: "), html.Span(element_types)]
                        ),
                        dbc.ListGroupItem(
                            [
                                html.Strong("Sample elements: "),
                                html.Span(str(sample_elements)),
                            ]
                        ),
                    ],
                    className="mb-3",
                ),
            ]
        )

    elif file_type == "DataFrame":
        # Information for DataFrame
        shape = file_info.get("Shape", "Unknown")
        columns = file_info.get("Columns", "Unknown")

        components.extend(
            [
                html.H6("DataFrame Details", className="mt-3"),
                dbc.ListGroup(
                    [
                        dbc.ListGroupItem([html.Strong("Shape: "), html.Span(shape)]),
                        dbc.ListGroupItem(
                            [html.Strong("Columns: "), html.Span(str(columns))]
                        ),
                    ],
                    className="mb-3",
                ),
            ]
        )

        # Try to parse and display a sample of the dataframe
        if "Sample data" in file_info:
            components.append(html.H6("Sample Data (First 5 rows):", className="mt-3"))
            try:
                if isinstance(file_info["Sample data"], dict):
                    sample_data = file_info["Sample data"]
                else:
                    # Try to parse from string (simplified)
                    sample_text = file_info["Sample data"]
                    if sample_text.startswith("{") and sample_text.endswith("}"):
                        sample_data = json.loads(sample_text.replace("'", '"'))
                    else:
                        raise ValueError("Cannot parse sample data")

                # Create a table from the sample data
                table_header = [html.Tr([html.Th(col) for col in sample_data.keys()])]

                # Get the number of rows in the first column
                first_col = list(sample_data.values())[0]
                num_rows = len(first_col) if isinstance(first_col, list) else 1

                # Create the table rows
                table_body = []
                for i in range(num_rows):
                    table_body.append(
                        html.Tr(
                            [
                                html.Td(
                                    sample_data[col][i]
                                    if isinstance(sample_data[col], list)
                                    and i < len(sample_data[col])
                                    else ""
                                )
                                for col in sample_data.keys()
                            ]
                        )
                    )

                components.append(
                    dbc.Table(
                        [html.Thead(table_header), html.Tbody(table_body)],
                        bordered=True,
                        hover=True,
                        responsive=True,
                    )
                )
            except Exception:
                components.append(
                    html.P("Failed to parse sample data", className="text-muted")
                )

    elif file_type == "ndarray":
        # Information for NumPy array
        shape = file_info.get("Shape", "Unknown")
        dtype = file_info.get("Data type", "Unknown")
        sample_data = file_info.get("Sample data", "Unknown")

        components.extend(
            [
                html.H6("NumPy Array Details", className="mt-3"),
                dbc.ListGroup(
                    [
                        dbc.ListGroupItem([html.Strong("Shape: "), html.Span(shape)]),
                        dbc.ListGroupItem(
                            [html.Strong("Data type: "), html.Span(dtype)]
                        ),
                        dbc.ListGroupItem(
                            [html.Strong("Sample data: "), html.Span(str(sample_data))]
                        ),
                    ],
                    className="mb-3",
                ),
            ]
        )

    elif file_type == "JSON" or file_type == "GeoJSON":
        # Information for JSON/GeoJSON files
        structure = file_info.get("Structure", "Unknown")

        # Create list items with basic info
        json_info_items = []

        if structure == "Dictionary":
            # For dictionary structure
            keys_count = file_info.get("Number of keys", "Unknown")
            top_keys = file_info.get("Top-level keys", [])

            json_info_items.extend(
                [
                    dbc.ListGroupItem(
                        [html.Strong("Structure: "), html.Span(structure)]
                    ),
                    dbc.ListGroupItem(
                        [html.Strong("Number of keys: "), html.Span(str(keys_count))]
                    ),
                    dbc.ListGroupItem(
                        [
                            html.Strong("Top-level keys: "),
                            html.Span(", ".join(str(k) for k in top_keys)),
                        ]
                    ),
                ]
            )

            # Add GeoJSON specific information if available
            if "GeoJSON Type" in file_info:
                json_type = file_info.get("GeoJSON Type", "")
                json_info_items.append(
                    dbc.ListGroupItem(
                        [html.Strong("GeoJSON Type: "), html.Span(json_type)]
                    )
                )

                if "Features Count" in file_info:
                    features_count = file_info.get("Features Count", 0)
                    json_info_items.append(
                        dbc.ListGroupItem(
                            [
                                html.Strong("Features Count: "),
                                html.Span(str(features_count)),
                            ]
                        )
                    )

                if "Geometry Type" in file_info:
                    geometry_type = file_info.get("Geometry Type", "Unknown")
                    json_info_items.append(
                        dbc.ListGroupItem(
                            [html.Strong("Geometry Type: "), html.Span(geometry_type)]
                        )
                    )

                if "Property Keys" in file_info:
                    property_keys = file_info.get("Property Keys", [])
                    if isinstance(property_keys, list):
                        property_keys_str = ", ".join(
                            str(k) for k in property_keys[:10]
                        )
                        if len(property_keys) > 10:
                            property_keys_str += "... (and more)"
                        json_info_items.append(
                            dbc.ListGroupItem(
                                [
                                    html.Strong("Property Keys: "),
                                    html.Span(property_keys_str),
                                ]
                            )
                        )
                    else:
                        json_info_items.append(
                            dbc.ListGroupItem(
                                [
                                    html.Strong("Property Keys: "),
                                    html.Span(str(property_keys)),
                                ]
                            )
                        )

        elif structure == "List":
            # For list structure
            items_count = file_info.get("Number of items", 0)
            item_type = file_info.get("Item Type", "Unknown")

            json_info_items.extend(
                [
                    dbc.ListGroupItem(
                        [html.Strong("Structure: "), html.Span(structure)]
                    ),
                    dbc.ListGroupItem(
                        [html.Strong("Number of items: "), html.Span(str(items_count))]
                    ),
                    dbc.ListGroupItem(
                        [html.Strong("Item Type: "), html.Span(item_type)]
                    ),
                ]
            )

            if item_type == "Dictionary" and "Sample Item Keys" in file_info:
                sample_keys = file_info.get("Sample Item Keys", [])
                sample_keys_str = ", ".join(str(k) for k in sample_keys[:10])
                if len(sample_keys) > 10:
                    sample_keys_str += "... (and more)"
                json_info_items.append(
                    dbc.ListGroupItem(
                        [html.Strong("Sample Item Keys: "), html.Span(sample_keys_str)]
                    )
                )

        components.extend(
            [
                html.H6(f"{file_type} Details", className="mt-3"),
                dbc.ListGroup(json_info_items, className="mb-3"),
            ]
        )

        # If we have a DataFrame representation, show it
        df_key = (
            "Properties DataFrame"
            if file_type == "GeoJSON"
            else "DataFrame Representation"
        )
        if df_key in file_info:
            df_info = file_info[df_key]

            components.extend(
                [
                    html.H6(f"{file_type} as DataFrame", className="mt-3"),
                    dbc.ListGroup(
                        [
                            dbc.ListGroupItem(
                                [
                                    html.Strong("Shape: "),
                                    html.Span(df_info.get("Shape", "Unknown")),
                                ]
                            ),
                            dbc.ListGroupItem(
                                [
                                    html.Strong("Columns: "),
                                    html.Span(
                                        ", ".join(
                                            str(c)
                                            for c in df_info.get("Columns", [])[:10]
                                        )
                                    ),
                                ]
                            ),
                        ],
                        className="mb-3",
                    ),
                ]
            )

            # Show sample data if available
            if "Sample Data" in df_info:
                components.append(html.H6("Sample Data:", className="mt-3"))
                sample_data = df_info.get("Sample Data", [])

                if (
                    sample_data
                    and isinstance(sample_data, list)
                    and len(sample_data) > 0
                ):
                    # Create a table from the sample data
                    if isinstance(sample_data[0], dict):
                        columns = list(sample_data[0].keys())

                        # Table header
                        table_header = [html.Tr([html.Th(col) for col in columns])]

                        # Table rows
                        table_body = []
                        for row in sample_data:
                            table_body.append(
                                html.Tr(
                                    [html.Td(str(row.get(col, ""))) for col in columns]
                                )
                            )

                        components.append(
                            dbc.Table(
                                [html.Thead(table_header), html.Tbody(table_body)],
                                bordered=True,
                                hover=True,
                                responsive=True,
                                size="sm",
                                className="mb-3",
                            )
                        )
                    else:
                        components.append(
                            html.Pre(
                                str(sample_data),
                                style={"maxHeight": "200px", "overflowY": "auto"},
                            )
                        )

    # Add collapsible card for the raw output
    unique_id = hash(f"{file_type}{str(hash(str(file_info)))}")
    components.append(
        html.Div(
            [
                dbc.Button(
                    "Show Raw Data",
                    id={"type": "collapse-button", "index": unique_id},
                    className="mb-3",
                    color="secondary",
                    size="sm",
                ),
                dbc.Collapse(
                    dbc.Card(
                        dbc.CardBody(
                            html.Pre(
                                json.dumps(file_info, indent=2),
                                style={"maxHeight": "300px", "overflowY": "auto"},
                            )
                        )
                    ),
                    id={"type": "collapse", "index": unique_id},
                    is_open=False,
                ),
            ]
        )
    )

    return html.Div(components)


@callback(
    Output({"type": "collapse", "index": dash.dependencies.MATCH}, "is_open"),
    [Input({"type": "collapse-button", "index": dash.dependencies.MATCH}, "n_clicks")],
    [State({"type": "collapse", "index": dash.dependencies.MATCH}, "is_open")],
)
def toggle_collapse(n_clicks, is_open):
    if n_clicks:
        return not is_open
    return is_open


@callback(
    Output({"type": "object-collapse", "index": dash.dependencies.MATCH}, "is_open"),
    [Input({"type": "object-button", "index": dash.dependencies.MATCH}, "n_clicks")],
    [State({"type": "object-collapse", "index": dash.dependencies.MATCH}, "is_open")],
)
def toggle_object_collapse(n_clicks, is_open):
    if n_clicks:
        return not is_open
    return is_open


# Define file type based on extension
def determine_file_type(filename):
    """
    Determine if a file is data or context based on its extension.

    Args:
        filename: The name of the file

    Returns:
        str: "data" or "context"
    """
    # Get the file extension
    ext = os.path.splitext(filename)[1].lower()

    # Define data file extensions
    data_extensions = [".pkl", ".pickle", ".csv", ".parquet", ".json", ".geojson"]

    # Define context file extensions
    context_extensions = [".txt", ".pdf", ".docx", ".doc"]

    if ext in data_extensions:
        return "data"
    elif ext in context_extensions:
        return "context"
    else:
        # Default to data for unknown extensions
        return "data"


@callback(
    Output("data-output", "children"),
    Output("loading-output", "children"),
    Output("files-being-analyzed", "children"),
    Output("error-message", "children"),
    Output("data-store", "data"),
    Output("context-output", "children"),
    Output("context-store", "data"),
    Input("unified-upload", "contents"),
    Input("unified-upload", "filename"),
    State("data-store", "data"),
    State("context-store", "data"),
)
def update_unified_output(
    contents, filenames, stored_data_results, stored_context_results
):
    """
    Callback for unified file upload.
    Files are automatically categorized and analyzed upon upload.

    Returns:
        - Data output component
        - Loading output
        - Files being analyzed text
        - Error message
        - Stored data results
        - Context output component
        - Stored context results
    """
    if contents is None or filenames is None:
        # No file uploaded yet
        return (
            html.Div(),
            "",
            "",
            "",
            {},
            html.Div("Upload context files to see extracted text."),
            {},
        )

    # Initialize data stores if not exists
    if stored_data_results is None:
        stored_data_results = {}
    if stored_context_results is None:
        stored_context_results = {}

    # Check if the callback was triggered by a file upload
    ctx = dash.callback_context
    if not ctx.triggered or ctx.triggered[0]["prop_id"] == ".":
        return (
            html.Div(),
            "",
            "",
            "",
            stored_data_results,
            html.Div("Upload context files to see extracted text."),
            stored_context_results,
        )

    # Get the trigger
    changed_id = ctx.triggered[0]["prop_id"].split(".")[0]
    if changed_id != "unified-upload":
        return (
            html.Div(),
            "",
            "",
            "",
            stored_data_results,
            html.Div("Upload context files to see extracted text."),
            stored_context_results,
        )

    # Sort files by type
    data_contents = []
    data_filenames = []
    context_contents = []
    context_filenames = []

    for content, filename in zip(contents, filenames):
        file_type = determine_file_type(filename)
        if file_type == "data":
            data_contents.append(content)
            data_filenames.append(filename)
        else:  # context
            context_contents.append(content)
            context_filenames.append(filename)

    # Process data files
    data_output_components = []
    new_data_results = dict(
        stored_data_results
    )  # Make a copy to preserve existing results

    # Show the files being analyzed
    files_display = ""
    if data_filenames:
        files_display = dbc.Alert(
            [
                html.H5("Analyzing Data Files:", className="alert-heading"),
                html.Ul([html.Li(filename) for filename in data_filenames]),
            ],
            color="info",
        )

    for content, filename in zip(data_contents, data_filenames):
        if content is None:
            continue

        # Save file to temporary directory
        try:
            file_path = save_uploaded_files(content, filename, "data")

            # Analyze file
            result = analyze_data_file(file_path)

            if "error" in result:
                data_output_components.append(
                    dbc.Alert(
                        [
                            html.H5(f"Error analyzing {filename}"),
                            html.P(result["error"]),
                        ],
                        color="danger",
                        className="mb-3",
                    )
                )
                log.error(
                    "Error analyzing data file",
                    error=result["error"],
                    filename=filename,
                )
            else:
                # Store analysis results
                new_data_results[filename] = result

                # Format analysis results for display
                file_card = create_file_summary_card(filename, result)
                data_output_components.append(file_card)

        except Exception as e:
            data_output_components.append(
                dbc.Alert(
                    [
                        html.H5(f"Error processing {filename}"),
                        html.P(str(e)),
                    ],
                    color="danger",
                    className="mb-3",
                )
            )
            log.error(
                "Exception while processing data file", error=str(e), filename=filename
            )

    # Process context files
    context_output_components = []
    new_context_results = dict(
        stored_context_results
    )  # Make a copy to preserve existing results

    for content, filename in zip(context_contents, context_filenames):
        if content is None:
            continue

        # Save file to temporary directory
        try:
            file_path = save_uploaded_files(content, filename, "context")

            # Extract text from context file
            result = extract_text_from_file(file_path)
            text_content = result.get("text", "")

            # Store the extracted text
            new_context_results[filename] = text_content

            # Create a card to display file information and preview
            context_card = dbc.Card(
                [
                    dbc.CardHeader(
                        [
                            html.H5(
                                [
                                    html.I(className="fas fa-file-alt me-2"),
                                    filename,
                                ],
                                className="mb-0",
                            )
                        ]
                    ),
                    dbc.CardBody(
                        [
                            html.P("Text Preview:"),
                            dbc.Card(
                                dbc.CardBody(
                                    html.P(
                                        text_content[:500] + "..."
                                        if len(text_content) > 500
                                        else text_content
                                    )
                                ),
                                className="bg-light",
                            ),
                        ]
                    ),
                ],
                className="mb-3",
            )
            context_output_components.append(context_card)

        except Exception as e:
            context_output_components.append(
                dbc.Alert(
                    [
                        html.H5(f"Error processing {filename}"),
                        html.P(str(e)),
                    ],
                    color="danger",
                    className="mb-3",
                )
            )
            log.error(
                "Exception while processing context file",
                error=str(e),
                filename=filename,
            )

    # Format data output
    data_output = html.Div("No valid data files uploaded.")
    if data_output_components:
        data_output = dbc.Row(
            [dbc.Col(component, md=6, lg=4) for component in data_output_components]
        )

    # Format context output
    context_output = html.Div("Upload context files to see extracted text.")
    if context_output_components:
        context_output = dbc.Row(
            [dbc.Col(component, md=6) for component in context_output_components]
        )

    return (
        data_output,
        "",
        files_display,
        "",
        new_data_results,
        context_output,
        new_context_results,
    )


@callback(
    Output("data-output", "children", allow_duplicate=True),
    Output("loading-output", "children", allow_duplicate=True),
    Output("files-being-analyzed", "children", allow_duplicate=True),
    Output("error-message", "children", allow_duplicate=True),
    Output("data-store", "data", allow_duplicate=True),
    Input("unified-upload", "contents"),
    Input("unified-upload", "filename"),
    State("data-store", "data"),
    prevent_initial_call=True,
)
def update_data_output(data_contents, data_filenames, stored_data_results):
    """
    Callback for data file upload (compatibility).
    """
    if data_contents is None or data_filenames is None:
        # No file uploaded yet
        return html.Div(), "", "", "", {}

    # Initialize data store if not exists
    if stored_data_results is None:
        stored_data_results = {}

    # Filter to only data files
    data_contents_filtered = []
    data_filenames_filtered = []

    for content, filename in zip(data_contents, data_filenames):
        if determine_file_type(filename) == "data":
            data_contents_filtered.append(content)
            data_filenames_filtered.append(filename)

    # The rest of your existing function here...
    # ... (continue with original implementation using data_contents_filtered and data_filenames_filtered)

    # Placeholder return - this should be updated with actual implementation
    return html.Div(), "", "", "", stored_data_results


@callback(
    Output("context-output", "children", allow_duplicate=True),
    Output("context-store", "data", allow_duplicate=True),
    Input("unified-upload", "contents"),
    Input("unified-upload", "filename"),
    State("context-store", "data"),
    prevent_initial_call=True,
)
def update_context_output(context_contents, context_filenames, stored_context_results):
    """
    Callback for context file upload (compatibility).
    """
    if context_contents is None or context_filenames is None:
        # No file uploaded yet
        return html.Div("Upload context files to see extracted text."), {}

    # Initialize data store if not exists
    if stored_context_results is None:
        stored_context_results = {}

    # Filter to only context files
    context_contents_filtered = []
    context_filenames_filtered = []

    for content, filename in zip(context_contents, context_filenames):
        if determine_file_type(filename) == "context":
            context_contents_filtered.append(content)
            context_filenames_filtered.append(filename)

    # The rest of your existing function here...
    # ... (continue with original implementation using context_contents_filtered and context_filenames_filtered)

    # Placeholder return - this should be updated with actual implementation
    return html.Div(
        "Upload context files to see extracted text."
    ), stored_context_results


@callback(
    Output("correlation-output", "children"),
    Output("correlation-store", "data"),
    Output("correlation-loading-output", "children"),
    Input("analyze-correlations-button", "n_clicks"),
    State("data-store", "data"),
)
def update_correlation_analysis(n_clicks, data_results):
    """
    Callback for updating the correlation analysis.
    Returns:
        - Correlation output component
        - Correlation results data
        - Loading output (for the spinner)
    """
    # Check if callback was triggered
    if n_clicks is None or n_clicks == 0 or not data_results:
        log.debug(
            "Correlation analysis not triggered",
            n_clicks=n_clicks,
            has_data=bool(data_results),
        )
        return (
            html.Div(
                "Click 'Analyze Correlations' to find connections between your data files."
            ),
            {},
            "",
        )

    # Check if we have data results
    if not data_results or len(data_results) < 2:
        log.warning(
            "Not enough data files for correlation analysis UI",
            file_count=len(data_results) if data_results else 0,
        )
        return (
            dbc.Alert(
                "Upload at least two data files to analyze correlations.",
                color="warning",
            ),
            {},
            "",
        )

    try:
        log.info(
            "Starting correlation analysis UI update", data_file_count=len(data_results)
        )
        # Find correlations
        correlation_results = correlate_data_files(data_results)

        log.info(
            "Processing correlation results for UI display",
            keys=list(correlation_results.keys()),
            has_error="error" in correlation_results,
            has_message="message" in correlation_results,
        )

        if "error" in correlation_results:
            log.warning(
                "Error in correlation results", error=correlation_results["error"]
            )
            return dbc.Alert(correlation_results["error"], color="warning"), {}, ""

        # Create output components for correlations
        output_components = []
        # Track if any correlations were found
        found_correlations = False

        # Shared columns
        shared_cols = correlation_results.get("shared_columns", {})
        if shared_cols:
            found_correlations = True
            log.info("Processing shared columns for UI", count=len(shared_cols))
            shared_cols_items = []
            for col, files in shared_cols.items():
                shared_cols_items.append(
                    dbc.ListGroupItem(
                        [html.Strong(col), ": Found in ", html.Span(", ".join(files))]
                    )
                )

            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(html.H5("Shared Columns", className="mb-0")),
                        dbc.CardBody(
                            [dbc.ListGroup(shared_cols_items, className="mb-3")]
                        ),
                    ],
                    className="mb-4",
                )
            )

        # Potential join keys
        joins = correlation_results.get("potential_joins", [])
        if joins:
            found_correlations = True
            log.info("Processing potential join keys for UI", count=len(joins))
            joins_items = []
            for join in joins:
                joins_items.append(
                    dbc.ListGroupItem(
                        [
                            html.Strong(join["column"]),
                            ": Could join ",
                            html.Span(", ".join(join["files"])),
                        ]
                    )
                )

            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(
                            html.H5("Potential Join Keys", className="mb-0")
                        ),
                        dbc.CardBody([dbc.ListGroup(joins_items, className="mb-3")]),
                    ],
                    className="mb-4",
                )
            )

        # Similar sized files
        similar_sizes = correlation_results.get("similar_sizes", [])
        if similar_sizes:
            found_correlations = True
            log.info("Processing similar sized files for UI", count=len(similar_sizes))
            size_items = []
            for size_info in similar_sizes:
                size_items.append(
                    dbc.ListGroupItem(
                        [
                            html.Span(", ".join(size_info["files"])),
                            ": Row counts of ",
                            html.Span(
                                ", ".join(
                                    [str(count) for count in size_info["row_counts"]]
                                )
                            ),
                        ]
                    )
                )

            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(
                            html.H5("Files with Similar Sizes", className="mb-0")
                        ),
                        dbc.CardBody([dbc.ListGroup(size_items, className="mb-3")]),
                    ],
                    className="mb-4",
                )
            )

        # Common data types
        common_types = correlation_results.get("common_data_types", {})
        if common_types:
            found_correlations = True
            log.info("Processing common data types for UI", count=len(common_types))
            types_components = []

            for files_key, types in common_types.items():
                type_items = []
                for col, dtype in types.items():
                    type_items.append(
                        dbc.ListGroupItem([html.Strong(col), ": ", html.Span(dtype)])
                    )

                types_components.append(
                    dbc.Card(
                        [
                            dbc.CardHeader(html.H6(files_key, className="mb-0")),
                            dbc.CardBody([dbc.ListGroup(type_items)]),
                        ],
                        className="mb-3",
                    )
                )

            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(html.H5("Common Data Types", className="mb-0")),
                        dbc.CardBody([html.Div(types_components)]),
                    ],
                    className="mb-4",
                )
            )

        # LLM-suggested join keys (NEW SECTION)
        llm_joins = correlation_results.get("llm_suggested_joins", [])
        if llm_joins:
            found_correlations = True
            log.info(
                "Processing LLM-suggested join keys for UI", raw_count=len(llm_joins)
            )
            # Sort by confidence score in descending order
            llm_joins = sorted(
                llm_joins, key=lambda x: x.get("confidence", 0), reverse=True
            )
            log.debug(
                "Sorted LLM joins",
                first_confidence=llm_joins[0].get("confidence", 0)
                if llm_joins
                else "N/A",
            )

            llm_joins_items = []
            for join in llm_joins:
                # Format the confidence as a percentage
                confidence = join.get("confidence", 0)
                confidence_str = f"{confidence * 100:.0f}%" if confidence else "Unknown"

                # Log simplified information about this join
                log.debug(f"Processing LLM join with confidence {confidence_str}")

                # Create badge with appropriate color based on confidence
                if confidence >= 0.7:
                    confidence_badge = dbc.Badge(
                        confidence_str, color="success", className="me-1"
                    )
                elif confidence >= 0.4:
                    confidence_badge = dbc.Badge(
                        confidence_str, color="warning", className="me-1"
                    )
                else:
                    confidence_badge = dbc.Badge(
                        confidence_str, color="secondary", className="me-1"
                    )

                llm_joins_items.append(
                    dbc.ListGroupItem(
                        [
                            html.Strong(
                                f"{join.get('source_file', 'Unknown')}:{join.get('source_column', 'Unknown')}"
                            ),
                            " ⟷ ",
                            html.Strong(
                                f"{join.get('target_file', 'Unknown')}:{join.get('target_column', 'Unknown')}"
                            ),
                            html.Br(),
                            confidence_badge,
                            html.Span(
                                join.get("explanation", "No explanation provided"),
                                className="text-muted small",
                            ),
                        ]
                    )
                )

            log.info("Created LLM joins UI items", count=len(llm_joins_items))

            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(
                            [
                                html.H5("AI-Suggested Join Keys", className="mb-0"),
                                html.Small(
                                    "Based on column names and sample values analysis",
                                    className="text-muted",
                                ),
                            ]
                        ),
                        dbc.CardBody(
                            [dbc.ListGroup(llm_joins_items, className="mb-3")]
                        ),
                    ],
                    className="mb-4 border-primary",  # Highlight this card with a primary border
                )
            )
        elif (
            "llm_no_joins_found" in correlation_results
            and correlation_results["llm_no_joins_found"]
        ):
            # LLM was called but didn't find any potential join keys
            # Don't set found_correlations = True here as this isn't a positive correlation finding
            log.info("LLM found no join keys, displaying informational card")
            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(
                            [
                                html.H5("AI-Suggested Join Keys", className="mb-0"),
                                html.Small(
                                    "Based on column names and sample values analysis",
                                    className="text-muted",
                                ),
                            ]
                        ),
                        dbc.CardBody(
                            [
                                html.P(
                                    [
                                        html.I(
                                            className="fas fa-info-circle me-2 text-info"
                                        ),
                                        "No potential join keys were identified by the AI analysis. This could mean:",
                                    ],
                                    className="mb-2",
                                ),
                                html.Ul(
                                    [
                                        html.Li(
                                            "The datasets may not be directly relatable"
                                        ),
                                        html.Li(
                                            "The column names and sample values don't provide enough context"
                                        ),
                                        html.Li(
                                            "The relationships may be more complex than direct column matches"
                                        ),
                                    ],
                                    className="mb-3",
                                ),
                                html.P(
                                    "Consider examining the data manually or providing more context files to help identify potential relationships.",
                                    className="small text-muted",
                                ),
                            ]
                        ),
                    ],
                    className="mb-4 border-info",
                )
            )
        elif "llm_analysis_error" in correlation_results:
            # Show error if LLM analysis failed
            log.warning(
                "LLM analysis error found",
                error=correlation_results["llm_analysis_error"],
            )
            output_components.append(
                dbc.Alert(
                    [
                        html.H6("AI Join Analysis Error", className="alert-heading"),
                        html.P(correlation_results["llm_analysis_error"]),
                    ],
                    color="warning",
                    className="mb-4",
                )
            )

        # Display a message if no correlations of any kind were found
        if not found_correlations:
            log.warning(
                "No correlations found for display",
                correlation_keys=list(correlation_results.keys()),
            )
            output_components.append(
                dbc.Alert(
                    [
                        html.H5("No Correlations Found", className="alert-heading"),
                        html.P(
                            "No correlations or relationships were found between the uploaded files. This could be because:"
                        ),
                        html.Ul(
                            [
                                html.Li("The datasets are completely unrelated"),
                                html.Li(
                                    "The column names are different and don't share common patterns"
                                ),
                                html.Li(
                                    "The data structures are too dissimilar to identify relationships automatically"
                                ),
                            ]
                        ),
                        html.P(
                            "Try adding more context files or examining the data manually to identify potential relationships."
                        ),
                    ],
                    color="info",
                    className="mb-4",
                )
            )

        log.info(
            "Correlation UI update complete",
            component_count=len(output_components),
            found_correlations=found_correlations,
        )

        # Return with empty loading output to clear spinner
        return html.Div(output_components), correlation_results, ""

    except Exception as e:
        log.error("Error in correlation analysis", error=str(e))
        return (
            dbc.Alert(
                [
                    html.H5("Error in Correlation Analysis"),
                    html.P(str(e)),
                    html.Pre(traceback.format_exc()),
                ],
                color="danger",
            ),
            {},
            "",
        )


@callback(
    Output("ai-plan-output", "children"),
    Output("ai-plan-store", "data"),
    Output("plan-loading-output", "children"),
    Input("generate-plan-button", "n_clicks"),
    State("data-store", "data"),
    State("context-store", "data"),
    State("correlation-store", "data"),
)
def update_analysis_plan(n_clicks, data_results, context_results, correlation_results):
    """
    Callback for generating AI analysis plan.
    Returns:
        - AI plan output component
        - AI plan data
        - Loading output (for the spinner)
    """
    if n_clicks is None or n_clicks == 0 or not data_results:
        return (
            html.Div("Click 'Generate Analysis Plan' to create an AI-powered plan."),
            {},
            "",
        )

    # Check if we have data results
    if not data_results:
        return (
            dbc.Alert(
                "Upload data files first before generating an analysis plan.",
                color="warning",
            ),
            {},
            "",
        )

    try:
        # Generate analysis plan
        plan_results = generate_analysis_plan(
            data_results, context_results or {}, correlation_results or {}
        )

        if "error" in plan_results:
            return (
                dbc.Alert(
                    [
                        html.H5("Error Generating Analysis Plan"),
                        html.P(plan_results["error"]),
                        html.Pre(plan_results.get("traceback", "")),
                    ],
                    color="danger",
                ),
                {},
                "",
            )

        # Create output components for the plan
        output_components = []

        # Add timestamp
        if "timestamp" in plan_results:
            output_components.append(
                dbc.Alert(
                    [
                        html.Strong("Generated at: "),
                        html.Span(plan_results["timestamp"]),
                    ],
                    color="info",
                    className="mb-3",
                )
            )

        # Analysis plan
        if "analysis_plan" in plan_results:
            # Handle different formats of analysis_plan (string or list)
            if isinstance(plan_results["analysis_plan"], str):
                plan_text = plan_results["analysis_plan"]
                # Try to split on numbered items
                plan_items = re.split(r"\n\d+\.|\n\-", plan_text)
                if len(plan_items) > 1:
                    plan_items = [item.strip() for item in plan_items if item.strip()]
                    plan_content = html.Ol([html.Li(item) for item in plan_items])
                else:
                    plan_content = html.Pre(plan_text)
            elif isinstance(plan_results["analysis_plan"], list):
                plan_content = html.Ol(
                    [html.Li(item) for item in plan_results["analysis_plan"]]
                )
            else:
                plan_content = html.Pre(str(plan_results["analysis_plan"]))

            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(html.H5("Analysis Plan", className="mb-0")),
                        dbc.CardBody([plan_content]),
                    ],
                    className="mb-4",
                )
            )

        # Analysis code
        if "analysis_code" in plan_results:
            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(
                            html.H5("Generated Python Code", className="mb-0")
                        ),
                        dbc.CardBody(
                            [
                                html.Button(
                                    "Execute in E2B Sandbox",
                                    id="execute-code-button",
                                    className="btn btn-success mb-3",
                                ),
                                html.Pre(
                                    plan_results["analysis_code"],
                                    style={
                                        "backgroundColor": "#f8f9fa",
                                        "padding": "15px",
                                        "borderRadius": "5px",
                                        "maxHeight": "500px",
                                        "overflowY": "auto",
                                    },
                                ),
                            ]
                        ),
                    ],
                    className="mb-4",
                )
            )

        # Return with empty loading output to clear spinner
        return html.Div(output_components), plan_results, ""

    except Exception as e:
        log.error("Error generating analysis plan", error=str(e))
        return (
            dbc.Alert(
                [
                    html.H5("Error Generating Analysis Plan"),
                    html.P(str(e)),
                    html.Pre(traceback.format_exc()),
                ],
                color="danger",
            ),
            {},
            "",
        )


@callback(
    Output("ai-results-output", "children"),
    Input("execute-code-button", "n_clicks"),
    State("ai-plan-store", "data"),
    State("data-store", "data"),
)
def update_execution_results(n_clicks, plan_data, data_store):
    """
    Callback for executing the generated code in E2B sandbox.
    Returns:
        - AI results output component
    """
    if (
        n_clicks is None
        or n_clicks == 0
        or not plan_data
        or "analysis_code" not in plan_data
    ):
        return html.Div("Execute the generated code to see results here.")

    try:
        log.info("Executing analysis code in E2B sandbox")

        # Get the code to execute
        analysis_code = plan_data["analysis_code"]
        if isinstance(analysis_code, list):
            analysis_code = "\n".join(analysis_code)

        # Get data file paths
        data_paths = {}
        if data_store:
            for filename, file_info in data_store.items():
                if "Path" in file_info:
                    data_paths[filename] = file_info["Path"]

        # Execute the code using our new function
        execution_results = execute_code_with_data(analysis_code, data_paths)

        output_components = []

        # Add execution status alert
        status_color = (
            "success" if execution_results.get("status") == "success" else "danger"
        )
        output_components.append(
            dbc.Alert(
                [
                    html.H5(
                        f"Execution {'Successful' if status_color == 'success' else 'Failed'}",
                        className="alert-heading",
                    ),
                    html.P(
                        f"Executed at: {execution_results.get('timestamp', 'Unknown')}"
                    ),
                ],
                color=status_color,
                className="mb-3",
            )
        )

        # Display any error
        if execution_results.get("error"):
            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(
                            html.H5("Execution Error", className="text-danger mb-0")
                        ),
                        dbc.CardBody(
                            [
                                html.Pre(
                                    execution_results.get("error", "Unknown error"),
                                    style={
                                        "backgroundColor": "#f8f9fa",
                                        "padding": "10px",
                                        "borderRadius": "5px",
                                        "color": "#dc3545",
                                    },
                                ),
                                html.Hr(),
                                html.Strong("Traceback:"),
                                html.Pre(
                                    execution_results.get(
                                        "traceback", "No traceback available"
                                    ),
                                    style={
                                        "backgroundColor": "#f8f9fa",
                                        "padding": "10px",
                                        "borderRadius": "5px",
                                        "maxHeight": "300px",
                                        "overflowY": "auto",
                                        "fontSize": "12px",
                                        "color": "#6c757d",
                                    },
                                ),
                            ]
                        ),
                    ],
                    className="mb-4",
                    color="danger",
                    outline=True,
                )
            )

        # Display plots if any
        plots = execution_results.get("plots", [])
        if plots:
            plot_items = []

            for i, plot_data in enumerate(plots):
                plot_items.append(
                    dbc.Col(
                        dbc.Card(
                            [
                                dbc.CardHeader(f"Plot {i + 1}"),
                                dbc.CardBody(
                                    html.Img(
                                        src=f"data:image/png;base64,{plot_data}",
                                        style={"width": "100%"},
                                    )
                                ),
                            ],
                            className="mb-3",
                        ),
                        md=6,
                    )
                )

            output_components.append(
                html.Div(
                    [
                        html.H4("Generated Plots", className="mb-3"),
                        dbc.Row(plot_items),
                    ],
                    className="mb-4",
                )
            )

        # Display tables if any
        tables = execution_results.get("tables", [])
        if tables:
            table_items = []

            for i, table_data in enumerate(tables):
                try:
                    # Convert table data to HTML table
                    df = pd.DataFrame(table_data)
                    html_table = dbc.Table.from_dataframe(
                        df, striped=True, bordered=True, hover=True, responsive=True
                    )

                    table_items.append(
                        dbc.Card(
                            [
                                dbc.CardHeader(f"Table {i + 1}"),
                                dbc.CardBody(html_table),
                            ],
                            className="mb-3",
                        )
                    )
                except Exception as e:
                    log.error(f"Error converting table data to HTML: {str(e)}")

            if table_items:
                output_components.append(
                    html.Div(
                        [
                            html.H4("Generated Tables", className="mb-3"),
                            html.Div(table_items),
                        ],
                        className="mb-4",
                    )
                )

        # Display stdout
        if execution_results.get("stdout"):
            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(html.H5("Execution Output", className="mb-0")),
                        dbc.CardBody(
                            html.Pre(
                                execution_results.get("stdout", "No output"),
                                style={
                                    "backgroundColor": "#f8f9fa",
                                    "padding": "10px",
                                    "borderRadius": "5px",
                                    "maxHeight": "400px",
                                    "overflowY": "auto",
                                },
                            )
                        ),
                    ],
                    className="mb-4",
                )
            )

        # Display stderr if any
        if execution_results.get("stderr"):
            output_components.append(
                dbc.Card(
                    [
                        dbc.CardHeader(html.H5("Standard Error", className="mb-0")),
                        dbc.CardBody(
                            html.Pre(
                                execution_results.get("stderr", "No errors"),
                                style={
                                    "backgroundColor": "#f8f9fa",
                                    "padding": "10px",
                                    "borderRadius": "5px",
                                    "maxHeight": "300px",
                                    "overflowY": "auto",
                                    "color": "#dc3545",
                                },
                            )
                        ),
                    ],
                    className="mb-4",
                    color="warning",
                    outline=True,
                )
            )

        # Display the executed code
        output_components.append(
            dbc.Card(
                [
                    dbc.CardHeader(html.H5("Executed Code", className="mb-0")),
                    dbc.CardBody(
                        html.Pre(
                            analysis_code,
                            style={
                                "backgroundColor": "#f8f9fa",
                                "padding": "10px",
                                "borderRadius": "5px",
                                "maxHeight": "300px",
                                "overflowY": "auto",
                            },
                        )
                    ),
                ],
                className="mb-4",
            )
        )

        return html.Div(output_components)

    except Exception as e:
        log.error("Error executing code in sandbox", error=str(e))
        return dbc.Alert(
            [
                html.H5("Error Executing Code"),
                html.P(str(e)),
                html.Pre(traceback.format_exc()),
            ],
            color="danger",
        )


def analyze_data_file(file_path: str) -> Dict[str, Any]:
    """
    Analyze different types of data files.

    Args:
        file_path: Path to the data file

    Returns:
        Dictionary with analysis results
    """
    # Get file information
    try:
        file_stats = os.stat(file_path)
        file_size = file_stats.st_size
        file_name = os.path.basename(file_path)

        file_info: Dict[str, Any] = {
            "Path": file_path,
            "Name": file_name,
            "Size": f"{file_size / 1024:.2f} KB"
            if file_size < 1024 * 1024
            else f"{file_size / (1024 * 1024):.2f} MB",
        }

        # Determine file type by extension
        _, ext = os.path.splitext(file_path)
        ext = ext.lower()

        # Process based on file type
        if ext in [".pkl", ".pickle"]:
            try:
                # For pickle files, analyze securely in E2B sandbox
                log.info(
                    "Analyzing pickle file securely in E2B sandbox",
                    file_path=file_path,
                    file_name=file_name,
                    file_size=file_size,
                )

                # Debug: Log the analyze_pickle_files function to ensure it exists
                log.debug(
                    "Type of analyze_pickle_files:",
                    type=str(type(analyze_pickle_files)),
                )

                # Create a string buffer to capture the output of analyze_pickle_files
                output_buffer = io.StringIO()
                original_stdout = sys.stdout
                sys.stdout = output_buffer

                # Call the function from pickle_analyzer to analyze the pickle in a sandbox
                log.debug("About to call analyze_pickle_files", file_path=file_path)
                analyze_pickle_files(
                    [file_path], verbose=True
                )  # Set verbose to True for more output
                log.debug("Finished calling analyze_pickle_files")

                # Restore stdout and get the captured output
                sys.stdout = original_stdout
                analysis_output = output_buffer.getvalue()

                # Debug: Log the raw output
                log.debug(
                    "Raw output from analyze_pickle_files",
                    output_length=len(analysis_output),
                    output_preview=analysis_output[:500]
                    + ("..." if len(analysis_output) > 500 else ""),
                )

                # We're going to parse the output directly without trying to match filenames
                file_info["Type"] = "Pickle File (Securely Analyzed)"
                file_info["Analysis Method"] = "E2B Sandbox"

                # Parse the output and associate all objects with this file
                analyzed_objects = parse_sandbox_output(analysis_output)

                if analyzed_objects:
                    # Add the object information to our results
                    file_info["Analyzed Objects"] = analyzed_objects
                    file_info["Num Objects"] = len(analyzed_objects)

                    # Extract a summary of object types
                    object_type_summary: Dict[str, int] = {}
                    for obj in analyzed_objects:
                        obj_type = obj.get("Object Type", "Unknown")
                        if obj_type in object_type_summary:
                            object_type_summary[obj_type] += 1
                        else:
                            object_type_summary[obj_type] = 1
                    file_info["Object Type Summary"] = object_type_summary

                    log.info(
                        "Successfully analyzed pickle file",
                        file_name=file_name,
                        num_objects=len(analyzed_objects),
                        object_types=object_type_summary,
                    )
                else:
                    log.warning(
                        "No objects found in pickle file analysis",
                        file_name=file_name,
                        analysis_output_sample=analysis_output[:200] + "..."
                        if analysis_output
                        else "No output",
                    )
                    file_info["Warning"] = "No objects found in analysis"
                    file_info["Raw Output"] = analysis_output

                return file_info
            except Exception as e:
                log.error(
                    "Failed to analyze pickle file in sandbox",
                    error=str(e),
                    traceback=traceback.format_exc(),
                    file_path=file_path,
                )
                return {
                    "error": f"Failed to analyze pickle file securely: {str(e)}",
                    "Type": "Pickle File",
                    "Analysis Method": "Failed E2B Sandbox Analysis",
                    "Path": file_path,
                    "Name": file_name,
                    "Size": file_info["Size"],
                }

        elif ext == ".csv":
            try:
                df = pd.read_csv(file_path, low_memory=False)
                file_info.update(analyze_dataframe(df))
                file_info["Type"] = "DataFrame (CSV)"
                return file_info
            except Exception as e:
                log.error(
                    "Failed to analyze CSV file", error=str(e), file_path=file_path
                )
                return {"error": f"Failed to analyze CSV file: {str(e)}"}
        elif ext in {".xls", ".xlsx"}:
            try:
                df = pd.read_excel(file_path)
                file_info.update(analyze_dataframe(df))
                file_info["Type"] = "DataFrame (XLSX)"
                return file_info
            except Exception as e:
                log.error(
                    "Failed to analyze XLSX file", error=str(e), file_path=file_path
                )
                return {"error": f"Failed to analyze Excel file: {str(e)}"}

        elif ext == ".parquet":
            try:
                df = pd.read_parquet(file_path)
                file_info.update(analyze_dataframe(df))
                file_info["Type"] = "DataFrame (Parquet)"
                return file_info
            except Exception as e:
                log.error(
                    "Failed to analyze Parquet file", error=str(e), file_path=file_path
                )
                return {"error": f"Failed to analyze Parquet file: {str(e)}"}

        elif ext in [".json", ".geojson"]:
            try:
                with open(file_path, "r") as f:
                    json_data = json.load(f)

                # Basic information about the JSON file
                file_info["Type"] = "GeoJSON" if ext == ".geojson" else "JSON"

                # Analyze the structure of the JSON
                if isinstance(json_data, dict):
                    file_info["Structure"] = "Dictionary"
                    file_info["Number of keys"] = len(json_data)
                    file_info["Top-level keys"] = list(json_data.keys())[
                        :10
                    ]  # Limit to first 10 keys

                    # Check for GeoJSON structure
                    if ext == ".geojson" or "type" in json_data:
                        json_type = json_data.get("type", "")
                        file_info["GeoJSON Type"] = json_type

                        # Handle different GeoJSON types
                        if json_type == "FeatureCollection" and "features" in json_data:
                            features = json_data["features"]
                            file_info["Features Count"] = len(features)

                            # Analyze a sample feature
                            if features and len(features) > 0:
                                sample_feature = features[0]
                                if "geometry" in sample_feature:
                                    file_info["Geometry Type"] = sample_feature[
                                        "geometry"
                                    ].get("type", "Unknown")

                                if "properties" in sample_feature:
                                    props = sample_feature["properties"]
                                    file_info["Property Keys"] = (
                                        list(props.keys())
                                        if isinstance(props, dict)
                                        else "No properties"
                                    )

                        elif json_type in ["Feature"]:
                            if "geometry" in json_data:
                                file_info["Geometry Type"] = json_data["geometry"].get(
                                    "type", "Unknown"
                                )

                            if "properties" in json_data:
                                props = json_data["properties"]
                                file_info["Property Keys"] = (
                                    list(props.keys())
                                    if isinstance(props, dict)
                                    else "No properties"
                                )

                elif isinstance(json_data, list):
                    file_info["Structure"] = "List"
                    file_info["Number of items"] = len(json_data)

                    # Analyze the first few items
                    if json_data:
                        sample_item = json_data[0]
                        if isinstance(sample_item, dict):
                            file_info["Item Type"] = "Dictionary"
                            if len(json_data) > 0:
                                file_info["Sample Item Keys"] = list(sample_item.keys())
                        else:
                            file_info["Item Type"] = type(sample_item).__name__

                # If it's a structure that can be converted to a DataFrame
                try:
                    # For GeoJSON, extract properties as a DataFrame
                    if (
                        ext == ".geojson"
                        and isinstance(json_data, dict)
                        and "features" in json_data
                    ):
                        # Extract properties from features
                        properties_list = []
                        for feature in json_data["features"]:
                            if "properties" in feature and isinstance(
                                feature["properties"], dict
                            ):
                                properties_list.append(feature["properties"])

                        if properties_list:
                            df = pd.DataFrame(properties_list)
                            df_info = analyze_dataframe(df)
                            file_info["Properties DataFrame"] = df_info
                    # For regular JSON that's a list of dictionaries
                    elif (
                        isinstance(json_data, list)
                        and json_data
                        and isinstance(json_data[0], dict)
                    ):
                        df = pd.DataFrame(json_data)
                        df_info = analyze_dataframe(df)
                        file_info["DataFrame Representation"] = df_info
                except Exception as df_err:
                    log.warning(
                        "Failed to convert JSON to DataFrame",
                        error=str(df_err),
                        file_path=file_path,
                    )

                return file_info
            except Exception as e:
                log.error(
                    f"Failed to analyze {'GeoJSON' if ext == '.geojson' else 'JSON'} file",
                    error=str(e),
                    file_path=file_path,
                )
                return {
                    "error": f"Failed to analyze {'GeoJSON' if ext == '.geojson' else 'JSON'} file: {str(e)}"
                }

        else:
            return {"error": f"Unsupported file type: {ext}"}

    except Exception as e:
        log.error("Failed to analyze data file", error=str(e), file_path=file_path)
        return {"error": f"Failed to analyze data file: {str(e)}"}


def analyze_dataframe(df: pd.DataFrame) -> Dict[str, Any]:
    """
    Analyze a pandas DataFrame and return its properties.

    Args:
        df: Pandas DataFrame to analyze

    Returns:
        Dictionary with DataFrame properties
    """
    result: Dict[str, Any] = {}

    # Basic DataFrame info
    result["Shape"] = str(df.shape)
    result["Columns"] = list(df.columns)
    result["dtypes"] = {col: str(dtype) for col, dtype in df.dtypes.items()}

    # Missing values
    missing_values = df.isnull().sum()
    result["Missing Values"] = {
        col: int(missing) for col, missing in missing_values.items() if missing > 0
    }

    # Sample data - Convert DataFrame to JSON-serializable format
    try:
        # First convert any Timestamp or datetime objects to strings
        sample_df = df.head(5).copy()

        # Convert all datetime columns to strings
        for col in sample_df.select_dtypes(
            include=["datetime64", "datetime64[ns]"]
        ).columns:
            sample_df[col] = sample_df[col].astype(str)

        # For any remaining Timestamp objects or other non-serializable types
        sample_data = sample_df.to_dict(orient="records")

        # Custom conversion for any remaining non-serializable objects
        for row in sample_data:
            for key, value in row.items():
                # Check for pandas Timestamp
                if hasattr(value, "timestamp") and callable(
                    getattr(value, "timestamp")
                ):
                    row[key] = str(value)
                # Handle numpy types
                elif isinstance(value, (np.integer, np.floating)):
                    row[key] = (
                        float(value) if isinstance(value, np.floating) else int(value)
                    )
                elif isinstance(value, np.ndarray):
                    row[key] = value.tolist()
                # Handle any other non-serializable types
                elif not isinstance(
                    value, (str, int, float, bool, list, dict, type(None))
                ):
                    row[key] = str(value)

        # Store sample data with consistent key name
        result["Sample Data"] = sample_data
    except Exception as e:
        # If serialization still fails, provide a simplified version
        log.error(f"Error serializing sample data: {str(e)}")
        result["Sample Data"] = []

    # Column statistics
    column_stats: Dict[str, Dict[str, Any]] = {}

    for col in df.columns:
        stats: Dict[str, Any] = {}
        dtype = str(df[col].dtype)

        if pd.api.types.is_numeric_dtype(df[col]):
            # For numeric columns
            stats["mean"] = (
                float(df[col].mean()) if not pd.isna(df[col].mean()) else None
            )
            stats["std"] = float(df[col].std()) if not pd.isna(df[col].std()) else None
            stats["min"] = float(df[col].min()) if not pd.isna(df[col].min()) else None
            stats["max"] = float(df[col].max()) if not pd.isna(df[col].max()) else None
            stats["median"] = (
                float(df[col].median()) if not pd.isna(df[col].median()) else None
            )

            # Only add if not all values are NaN
            if any(v is not None for v in stats.values()):
                column_stats[str(col)] = stats

        elif pd.api.types.is_string_dtype(df[col]) or pd.api.types.is_object_dtype(
            df[col]
        ):
            # For string/object columns
            value_counts = df[col].value_counts(dropna=False).head(5).to_dict()

            # Convert any non-serializable keys or values
            value_counts_serializable = {}
            for k, v in value_counts.items():
                # Convert key if it's a Timestamp or other non-serializable type
                key = str(k) if not isinstance(k, (str, int, float, bool)) else k
                # Convert value to int (should be a count)
                val = int(v)
                value_counts_serializable[key] = val

            stats["unique_count"] = df[col].nunique()
            stats["top_values"] = value_counts_serializable

            column_stats[str(col)] = stats

    result["Column Stats"] = column_stats

    return result


def extract_text_from_file(file_path: str) -> Dict[str, Any]:
    """
    Extract text content from various file types.

    Args:
        file_path: Path to the context file

    Returns:
        Dictionary with extracted text and possibly error messages
    """
    file_name = os.path.basename(file_path)
    file_stats = os.stat(file_path)
    file_size = file_stats.st_size

    result: Dict[str, Any] = {
        "path": file_path,
        "name": file_name,
        "size": f"{file_size / 1024:.2f} KB"
        if file_size < 1024 * 1024
        else f"{file_size / (1024 * 1024):.2f} MB",
    }

    # Determine file type by extension
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()

    text_content = ""

    try:
        if ext == ".txt":
            with open(file_path, "r", encoding="utf-8") as f:
                text_content = f.read()
            result["type"] = "Text File"

        elif ext == ".pdf":
            try:
                import PyPDF2

                with open(file_path, "rb") as f:
                    pdf_reader = PyPDF2.PdfReader(f)
                    for page in pdf_reader.pages:
                        text_content += page.extract_text() + "\n\n"
                result["type"] = "PDF File"
            except ImportError:
                result["error"] = "PyPDF2 library not installed. Unable to process PDF."
                log.error("PyPDF2 library not installed")

        elif ext in [".docx", ".doc"]:
            try:
                import docx

                doc = docx.Document(file_path)
                for para in doc.paragraphs:
                    text_content += para.text + "\n"
                result["type"] = "Word Document"
            except ImportError:
                result["error"] = (
                    "python-docx library not installed. Unable to process Word document."
                )
                log.error("python-docx library not installed")
        else:
            result["error"] = f"Unsupported file type: {ext}"
            return result

        # Truncate text content if too large
        if len(text_content) > 1000000:  # ~1MB of text
            text_content = text_content[:1000000] + "... [truncated]"
            result["warning"] = "Text content was truncated due to large size"

        result["text"] = text_content
        # For backward compatibility
        result["Text"] = text_content
        return result

    except Exception as e:
        log.error("Failed to extract text from file", error=str(e), file_path=file_path)
        result["error"] = f"Failed to extract text: {str(e)}"
        return result


def correlate_data_files(data_results: Dict[str, Dict[str, Any]]) -> Dict[str, Any]:
    """
    Find correlations between data files

    Args:
        data_results: Dictionary of data file analysis results

    Returns:
        Dictionary with correlation results
    """
    log.info("Argument", arg=data_results)
    if not data_results or len(data_results) < 2:
        log.warning(
            "Not enough data files for correlation analysis",
            count=len(data_results) if data_results else 0,
        )
        return {"error": "Need at least two data files to find correlations"}

    correlations: Dict[str, Any] = {
        "shared_columns": {},
        "common_data_types": {},
        "similar_sizes": [],
        "potential_joins": [],
        "llm_suggested_joins": [],  # New field for LLM-identified join keys
    }

    # Find common columns across DataFrames
    dataframes = {
        fname: info
        for fname, info in data_results.items()
        if info.get("Type", "").startswith("DataFrame")
    }

    log.info(
        "DataFrame files identified for correlation",
        count=len(dataframes),
        filenames=list(dataframes.keys()),
    )

    # Exit early if we don't have at least 2 DataFrames
    if len(dataframes) < 2:
        log.warning(
            "Not enough DataFrame files for correlation analysis",
            dataframes_found=len(dataframes),
        )
        return {
            "message": "Not enough DataFrame files found for correlation analysis",
            "dataframes_found": len(dataframes),
        }

    # Track all columns seen
    all_columns: Dict[str, List[str]] = {}

    # Find shared columns
    for fname, info in dataframes.items():
        columns = info.get("Columns", [])
        log.debug(f"Processing columns for file {fname}", column_count=len(columns))
        for col in columns:
            col_str = str(col)
            if col_str not in all_columns:
                all_columns[col_str] = []
            all_columns[col_str].append(fname)

    # Filter for columns that appear in multiple files
    correlations["shared_columns"] = {
        col: files for col, files in all_columns.items() if len(files) > 1
    }
    log.info("Shared columns identified", count=len(correlations["shared_columns"]))

    # Find potential join keys
    potential_joins: List[Dict[str, Any]] = []
    for col, files in correlations["shared_columns"].items():
        if len(files) >= 2:
            potential_joins.append({"column": col, "files": files})
    correlations["potential_joins"] = potential_joins
    log.info("Potential join keys identified", count=len(potential_joins))

    # Find similar data types
    common_types_count = 0
    for fname1, info1 in dataframes.items():
        for fname2, info2 in dataframes.items():
            if fname1 >= fname2:  # Skip self-comparisons and duplicates
                continue

            common_cols = set(info1.get("Columns", [])).intersection(
                set(info2.get("Columns", []))
            )
            if common_cols:
                dtypes1 = info1.get("dtypes", {})
                dtypes2 = info2.get("dtypes", {})

                matching_dtypes = {}
                for col in common_cols:
                    col_str = str(col)
                    if col_str in dtypes1 and col_str in dtypes2:
                        if dtypes1[col_str] == dtypes2[col_str]:
                            matching_dtypes[col_str] = dtypes1[col_str]

                if matching_dtypes:
                    key = f"{fname1} ⟷ {fname2}"
                    if "common_data_types" not in correlations:
                        correlations["common_data_types"] = {}
                    correlations["common_data_types"][key] = matching_dtypes
                    common_types_count += 1

    log.info("Common data types identified", count=common_types_count)

    # Find files with similar sizes
    sizes = []
    for fname, info in dataframes.items():
        shape_str = info.get("Shape", "(0, 0)")
        # Handle shape stored as string like "(100, 5)"
        if (
            isinstance(shape_str, str)
            and shape_str.startswith("(")
            and shape_str.endswith(")")
        ):
            try:
                # Parse the row count from the shape string
                row_count = int(shape_str.strip("()").split(",")[0].strip())
                sizes.append((fname, row_count))
            except (ValueError, IndexError):
                continue
        elif isinstance(shape_str, tuple) and len(shape_str) == 2:
            # Handle shape stored as actual tuple
            sizes.append((fname, shape_str[0]))

    sizes.sort(key=lambda x: x[1])

    similar_sizes: List[Dict[str, Any]] = []
    for i in range(len(sizes) - 1):
        curr_file, curr_size = sizes[i]
        next_file, next_size = sizes[i + 1]

        # If within 10% of each other
        if curr_size > 0 and abs(curr_size - next_size) / curr_size < 0.1:
            similar_sizes.append(
                {"files": [curr_file, next_file], "row_counts": [curr_size, next_size]}
            )
    correlations["similar_sizes"] = similar_sizes
    log.info("Similar sized files identified", count=len(similar_sizes))

    # NEW PART: Use LLM to identify potential join keys based on column names and sample values
    if openai_client:
        log.info("Starting LLM-based join key analysis")
        try:
            # Prepare detailed information about each dataframe for the LLM
            dataframe_details = {}
            missing_sample_data = []

            for fname, info in dataframes.items():
                # Extract column information
                columns = info.get("Columns", [])

                # Extract sample data if available - check different possible keys
                sample_data = []

                # Try different keys that might contain sample data
                sample_data_found = False
                sample_key_used = None

                for key in ["Sample Data", "sample_data", "SampleData"]:
                    if key in info and info[key]:
                        raw_sample = info[key]
                        log.debug(
                            f"Found sample data in key '{key}' for file {fname}",
                            is_list=isinstance(raw_sample, list),
                            is_str=isinstance(raw_sample, str),
                        )

                        # Handle different formats
                        if isinstance(raw_sample, list):
                            sample_data = raw_sample
                            sample_data_found = True
                            sample_key_used = key
                            break
                        elif isinstance(raw_sample, str):
                            # Try to parse JSON string
                            try:
                                if raw_sample.startswith("["):
                                    sample_data = json.loads(
                                        raw_sample.replace("'", '"')
                                    )
                                    sample_data_found = True
                                    sample_key_used = key
                                    break
                                elif raw_sample.startswith("{"):
                                    # Might be a dict format with columns as keys
                                    sample_dict = json.loads(
                                        raw_sample.replace("'", '"')
                                    )
                                    # Convert to records format if possible
                                    if all(
                                        isinstance(sample_dict[col], list)
                                        for col in sample_dict
                                    ):
                                        # Get the length of the first list
                                        first_col = list(sample_dict.keys())[0]
                                        length = len(sample_dict[first_col])
                                        # Create records
                                        records = []
                                        for i in range(length):
                                            record = {}
                                            for col, values in sample_dict.items():
                                                if i < len(values):
                                                    record[col] = values[i]
                                            records.append(record)
                                        sample_data = records
                                        sample_data_found = True
                                        sample_key_used = key
                                        break
                            except Exception as e:
                                log.warning(
                                    f"Failed to parse sample data JSON for file {fname}",
                                    error=str(e),
                                    key=key,
                                )
                                # Failed to parse, continue to next key
                                pass

                if not sample_data_found:
                    missing_sample_data.append(fname)
                    log.warning(f"No sample data found for file {fname}")
                else:
                    log.info(
                        f"Sample data found for file {fname}",
                        key=sample_key_used,
                        records_count=len(sample_data),
                    )

                # Create a summary of sample values for each column
                column_samples = {}
                for col in columns:
                    col_str = str(col)
                    col_values = []

                    # Extract values from sample data
                    for row in sample_data[:3]:  # Use up to 3 rows
                        if isinstance(row, dict) and col_str in row:
                            value = row[col_str]
                            # Convert non-string values to strings
                            if not isinstance(value, str):
                                value = str(value)
                            col_values.append(value)

                    if col_values:
                        column_samples[col_str] = col_values

                # Add to dataframe details
                dataframe_details[fname] = {
                    "columns": columns,
                    "column_samples": column_samples,
                    "shape": info.get("Shape", "Unknown"),
                    "dtypes": info.get("dtypes", {}),
                }

            log.info(
                "Prepared dataframe details for LLM",
                dataframe_count=len(dataframe_details),
                missing_sample_data=missing_sample_data,
            )

            # Prepare the prompt for the LLM
            prompt = f"""
You are analyzing multiple dataframes to identify potential join keys. I'll provide details about each dataframe including column names and sample values.

Your task is to identify columns across different dataframes that could potentially be used as join keys, even if they have different column names. 
Look for columns that might contain the same type of entity identifiers, foreign keys, or related information.

Consider:
1. Columns with similar names or semantically related names (e.g. 'user_id' and 'uid')
2. Columns with similar value patterns
3. Columns that might represent the same entities across different tables

Here are the dataframes:
{json.dumps(dataframe_details, indent=2)}

Respond with a JSON array where each item represents a potential join relationship. Each item should be an object with these fields:
- "source_file": String, name of the first dataframe
- "source_column": String, column name in the first dataframe
- "target_file": String, name of the second dataframe
- "target_column": String, column name in the second dataframe
- "confidence": Number between 0 and 1, your confidence in this relationship
- "explanation": String, brief explanation of why these columns might be joinable

Only include column pairs that have a reasonable chance of being related. Focus on quality over quantity.
If you're not confident about any join keys, return an empty array [].
"""

            log.info(
                "Sending prompt to OpenAI for join key analysis",
                prompt_length=len(prompt),
                dataframe_count=len(dataframe_details),
            )

            # Call the OpenAI API
            response = openai_client.chat.completions.create(
                model="gpt-4o",
                messages=[
                    {
                        "role": "system",
                        "content": "You are a data analysis expert that identifies potential relationships between dataframes.",
                    },
                    {"role": "user", "content": prompt},
                ],
                response_format={"type": "json_object"},
                temperature=0.3,  # Lower temperature for more conservative suggestions
            )

            # Extract and parse the response
            llm_response = response.choices[0].message.content
            if llm_response:
                log.info(
                    "Received response from OpenAI",
                    response_length=len(llm_response),
                    response_snippet=llm_response[:100],
                )

            # Log the full response for debugging
            log.debug("Full LLM response", response=llm_response)

            # Parse the JSON response
            if llm_response:
                try:
                    suggested_joins = json.loads(llm_response)
                    # Log the raw parsed JSON
                    log.debug(
                        "Parsed LLM response JSON",
                        is_dict=isinstance(suggested_joins, dict),
                        is_list=isinstance(suggested_joins, list),
                    )

                    # Check if it's a list or if it's wrapped in another object
                    if isinstance(suggested_joins, dict) and "joins" in suggested_joins:
                        suggested_joins = suggested_joins["joins"]
                        log.info("Extracted 'joins' array from response dictionary")
                    elif isinstance(suggested_joins, list):
                        # Already in the right format
                        log.info("LLM response is already in list format")
                    elif isinstance(suggested_joins, dict):
                        # Look for any array in the response
                        found_array = False
                        for key, value in suggested_joins.items():
                            if isinstance(value, list):
                                log.info(
                                    f"Found array in key '{key}' of response dictionary"
                                )
                                suggested_joins = value
                                found_array = True
                                break
                        if not found_array:
                            # No array found, create empty list
                            log.warning(
                                "No array found in response dictionary, using empty list"
                            )
                            suggested_joins = []
                    else:
                        # Not a recognized format, create empty list
                        log.warning(
                            "Unrecognized format in LLM response, using empty list",
                            type=type(suggested_joins).__name__,
                        )
                        suggested_joins = []

                    # Ensure it's a list
                    if not isinstance(suggested_joins, list):
                        log.warning(
                            "Ensuring suggested_joins is a list",
                            original_type=type(suggested_joins).__name__,
                        )
                        suggested_joins = []

                    correlations["llm_suggested_joins"] = suggested_joins
                    # Add a flag to indicate the LLM was successfully called but found no joins
                    if len(suggested_joins) == 0:
                        log.info("LLM found no join keys")
                        correlations["llm_no_joins_found"] = True

                    log.info(
                        "LLM suggested join keys analysis complete",
                        count=len(suggested_joins),
                        keys=[
                            (
                                j.get("source_file", "?"),
                                j.get("source_column", "?"),
                                j.get("target_file", "?"),
                                j.get("target_column", "?"),
                            )
                            for j in suggested_joins[:5]
                        ],
                    )  # Log first 5 joins
                except json.JSONDecodeError as e:
                    log.error(
                        "Failed to parse LLM response as JSON",
                        error=str(e),
                        response=llm_response[:200],
                    )
                    correlations["llm_analysis_error"] = f"JSON parse error: {str(e)}"
        except Exception as e:
            log.error(
                "Error in LLM-based join key analysis",
                error=str(e),
                traceback=traceback.format_exc(),
            )
            # Don't fail the whole function if LLM analysis fails
            correlations["llm_analysis_error"] = str(e)
    else:
        log.warning("OpenAI client not available, skipping LLM-based join key analysis")

    log.info(
        "Correlation analysis complete",
        shared_columns=len(correlations.get("shared_columns", {})),
        potential_joins=len(correlations.get("potential_joins", [])),
        similar_sizes=len(correlations.get("similar_sizes", [])),
        common_data_types=len(correlations.get("common_data_types", {})),
        llm_suggested_joins=len(correlations.get("llm_suggested_joins", [])),
    )

    return correlations


def generate_analysis_plan(
    data_results: Dict[str, Dict[str, Any]],
    context_texts: Dict[str, str],
    correlations: Dict[str, Any],
) -> Dict[str, Any]:
    """
    Generate an analysis plan using OpenAI API

    Args:
        data_results: Dictionary of data file analysis results
        context_texts: Dictionary of context file texts
        correlations: Dictionary of correlation results

    Returns:
        Dictionary with analysis plan and generated code
    """
    if not openai_client:
        return {"error": "OpenAI client not initialized. Please check your API key."}

    # Prepare the prompt
    data_summary = json.dumps(
        {
            k: {
                "type": v.get("Type", "Unknown"),
                "shape": v.get("Shape", "Unknown"),
                "columns": v.get("Columns", [])[:10],  # Limit columns for prompt size
            }
            for k, v in data_results.items()
        },
        indent=2,
    )

    # Prepare context text (truncate if too long)
    context_summary = ""
    for filename, text in context_texts.items():
        # Truncate text if too long
        if len(text) > 500:
            text = text[:500] + "... [truncated]"
        context_summary += f"=== {filename} ===\n{text}\n\n"

    correlation_summary = json.dumps(correlations, indent=2)

    prompt = f"""
As a data science assistant, analyze the following datasets and context information to create an analysis plan.

DATA FILES:
{data_summary}

CORRELATIONS BETWEEN DATA FILES:
{correlation_summary}

CONTEXT FROM TEXT FILES:
{context_summary}

Based on this information, please:
1. Create a concise analysis plan
2. Generate Python code that could be executed in an E2B sandbox to analyze these datasets
3. The code should focus on finding patterns, correlations, and insights between the datasets
4. Utilize any context information to guide the analysis
5. Include visualizations where appropriate

Respond with a JSON object containing:
- "analysis_plan": A step-by-step plan for data analysis
- "analysis_code": Complete Python code that could be executed

In "analysis_code":

- The code should be self-contained and handle reading the files from their paths.
- It should contain no placeholders or truncations, only valid Python.
- Avoid triple backticks at the start or the end, just issue Python straightaway.
- Do *NOT* use geopandas.
"""

    try:
        response = openai_client.chat.completions.create(
            model="gpt-4o",
            seed=SEED,
            messages=[
                {
                    "role": "system",
                    "content": "You are a data science assistant that generates analysis plans and Python code.",
                },
                {"role": "user", "content": prompt},
            ],
            response_format={"type": "json_object"},
        )

        # Extract the response
        response_text = response.choices[0].message.content
        if response_text:
            result = json.loads(response_text)

            # Add timestamp
            result["timestamp"] = pd.Timestamp.now().isoformat()

            return result
        else:
            return {"error": "No response received from OpenAI API"}

    except Exception as e:
        return {
            "error": f"Failed to generate analysis plan: {str(e)}",
            "traceback": traceback.format_exc(),
        }


def parse_sandbox_output(analysis_text: str) -> List[Dict[str, Any]]:
    """
    Parse the raw output from the E2B sandbox analysis of pickle files.

    Instead of trying to associate each section with a filename, we treat each section
    as an individual object analysis and return a list of analyzed objects.

    Args:
        analysis_text: The raw text output from pickle_analyzer

    Returns:
        List of dictionaries with analysis results for each object
    """
    # Debug log the analysis text
    log.debug(
        "Parsing sandbox output",
        text_length=len(analysis_text),
        text_sample=analysis_text[:200] + ("..." if len(analysis_text) > 200 else ""),
    )

    # Split into sections by the separator (50 equal signs)
    sections = analysis_text.split("=" * 50)
    sections = [s.strip() for s in sections if s.strip()]

    log.debug(f"Found {len(sections)} sections in analysis text")

    analyzed_objects = []

    for i, section in enumerate(sections):
        lines = section.split("\n")
        if len(lines) < 2:
            log.warning(f"Section {i} too short to parse", lines_count=len(lines))
            continue

        # Extract information about this object
        object_info = {}
        current_key = None
        current_value = []

        # Try to extract object type directly from the first few lines
        object_type = "Unknown"
        for idx, line in enumerate(lines[:10]):  # Check first 10 lines
            if "Object Type:" in line:
                try:
                    object_type = line.split("Object Type:")[1].strip()
                    break
                except:
                    pass
            elif "Type:" in line:
                try:
                    object_type = line.split("Type:")[1].strip()
                    break
                except:
                    pass

        # Check for specific types in output
        if object_type == "Unknown":
            # Look for specific patterns in text
            if any("dict" in line.lower() for line in lines[:20]):
                object_type = "dict"
            elif any("list" in line.lower() for line in lines[:20]):
                object_type = "list"
            elif any("dataframe" in line.lower() for line in lines[:20]):
                object_type = "DataFrame"
            elif any("ndarray" in line.lower() for line in lines[:20]):
                object_type = "ndarray"

        # Explicitly set the Object Type
        object_info["Object Type"] = object_type

        # Check if this section contains object analysis
        contains_analysis = False
        for line in lines:
            if (
                "Object Type:" in line
                or "Type:" in line
                or "Value:" in line
                or "Contents:" in line
                or "Length:" in line
                or "Shape:" in line
                or "Size:" in line
                or "keys_count" in line
                or "Number of keys:" in line
            ):
                contains_analysis = True
                break

        if not contains_analysis:
            log.debug(
                f"Section {i} doesn't appear to contain object analysis, skipping"
            )
            continue

        # Process line by line
        for line in lines:
            line = line.strip()
            if not line:
                continue

            # If it's a new key-value pair
            if ": " in line and not line.startswith(" "):
                # Save the previous key-value pair if it exists
                if current_key:
                    object_info[current_key] = (
                        "\n".join(current_value)
                        if len(current_value) > 1
                        else current_value[0]
                    )
                    current_value = []

                # Start a new key-value pair
                try:
                    current_key, value = line.split(": ", 1)

                    # Normalize key names for consistency with UI expectations
                    if current_key.lower() == "type":
                        current_key = "Object Type"
                        # Update the object_type variable as well
                        object_type = value
                    elif current_key.lower() == "size":
                        current_key = "Size (bytes)"
                    elif current_key.lower() == "length":
                        current_key = "Length"
                    elif current_key.lower() == "shape":
                        current_key = "Shape"
                    elif current_key.lower() == "number of keys":
                        current_key = "Number of Keys"

                    current_value.append(value)
                except ValueError:
                    # If there's a problem with splitting, just store the whole line
                    log.warning(f"Could not parse line properly: {line}")
                    if current_key:
                        current_value.append(line)
            else:
                # Continue with the current value
                if current_key:
                    current_value.append(line)

        # Save the last key-value pair
        if current_key:
            object_info[current_key] = (
                "\n".join(current_value) if len(current_value) > 1 else current_value[0]
            )

        # Make sure Object Type is set
        if "Object Type" not in object_info:
            object_info["Object Type"] = object_type

        # Add raw section text for debugging
        object_info["Raw Section Text"] = (
            section[:500] + "..." if len(section) > 500 else section
        )

        if object_info:
            # Add an index to identify the object
            object_info["Object Index"] = str(i)
            analyzed_objects.append(object_info)
            log.debug(f"Added object #{i} with {len(object_info)} properties")

    # Calculate object type summary
    object_type_summary: Dict[str, int] = {}
    for obj in analyzed_objects:
        obj_type = obj.get("Object Type", "Unknown")
        if obj_type in object_type_summary:
            object_type_summary[obj_type] += 1
        else:
            object_type_summary[obj_type] = 1

    log.info(f"Finished parsing {len(analyzed_objects)} objects from sandbox output")
    log.info(f"Object type summary: {object_type_summary}")

    return analyzed_objects


def main():
    """Run the Dash app server"""
    app.run(debug=True)


if __name__ == "__main__":
    main()
